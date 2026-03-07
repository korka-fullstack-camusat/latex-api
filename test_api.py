"""
Test automatisé — Word to LaTeX API
Usage: python3 test_api.py [--url http://localhost:8000]
"""
import argparse
import io
import os
import sys
import zipfile
from dataclasses import dataclass, field
from typing import Callable

import requests
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── Config ────────────────────────────────────────────────────────────────────
BASE_URL = "http://localhost:8000"
CONVERT   = "/convert"
HEALTH    = "/health"

GREEN = "\033[92m"
RED   = "\033[91m"
YELLOW= "\033[93m"
BOLD  = "\033[1m"
RESET = "\033[0m"


# ── Result helpers ─────────────────────────────────────────────────────────────
@dataclass
class TestResult:
    name: str
    passed: bool
    detail: str = ""
    warnings: list[str] = field(default_factory=list)


results: list[TestResult] = []


def ok(name: str, detail: str = "") -> TestResult:
    r = TestResult(name, True, detail)
    results.append(r)
    print(f"  {GREEN}✔{RESET} {name}" + (f"  → {detail}" if detail else ""))
    return r


def fail(name: str, detail: str = "") -> TestResult:
    r = TestResult(name, False, detail)
    results.append(r)
    print(f"  {RED}✘{RESET} {name}" + (f"  → {detail}" if detail else ""))
    return r


def warn(result: TestResult, msg: str):
    result.warnings.append(msg)
    print(f"    {YELLOW}⚠{RESET} {msg}")


def section(title: str):
    print(f"\n{BOLD}{'─'*55}{RESET}")
    print(f"{BOLD}  {title}{RESET}")
    print(f"{BOLD}{'─'*55}{RESET}")


# ── .docx factories ───────────────────────────────────────────────────────────

def make_simple_doc() -> bytes:
    """Document texte basique."""
    doc = Document()
    doc.add_heading("Introduction au Machine Learning", 0)
    doc.add_heading("Définition", 1)
    doc.add_paragraph(
        "Le machine learning est une branche de l'intelligence artificielle "
        "qui permet aux systèmes d'apprendre automatiquement à partir des données."
    )
    doc.add_heading("Types d'apprentissage", 2)
    doc.add_paragraph("Il existe trois grandes catégories :")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_rich_formatting_doc() -> bytes:
    """Gras, italique, souligné, exposant, indice."""
    doc = Document()
    doc.add_heading("Test de Formatage", 1)

    p = doc.add_paragraph()
    p.add_run("Texte normal, ").bold = False
    r = p.add_run("texte gras, ")
    r.bold = True
    r = p.add_run("texte italique, ")
    r.italic = True
    r = p.add_run("texte souligné")
    r.underline = True

    p2 = doc.add_paragraph()
    p2.add_run("H")
    r = p2.add_run("2")
    r.font.subscript = True
    p2.add_run("O et E=mc")
    r2 = p2.add_run("2")
    r2.font.superscript = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_lists_doc() -> bytes:
    """Listes à puces et numérotées."""
    doc = Document()
    doc.add_heading("Avantages du LaTeX", 1)
    doc.add_paragraph("Avantages :", style="List Bullet")
    for item in ["Qualité typographique", "Gestion des références", "Portabilité"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Étapes d'installation", 2)
    for i, step in enumerate([
        "Télécharger TeX Live",
        "Lancer l'installateur",
        "Configurer le PATH",
        "Vérifier avec pdflatex --version",
    ], 1):
        doc.add_paragraph(f"{i}. {step}", style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_table_doc() -> bytes:
    """Tableau de données."""
    doc = Document()
    doc.add_heading("Comparaison des modèles", 1)

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    headers = ["Modèle", "Précision", "Vitesse (ms)"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True

    data = [
        ("ResNet-50", "92.3%", "45"),
        ("EfficientNet", "94.1%", "38"),
        ("ViT-Base", "95.6%", "120"),
    ]
    for row_idx, (model, acc, speed) in enumerate(data, 1):
        row = table.rows[row_idx]
        row.cells[0].text = model
        row.cells[1].text = acc
        row.cells[2].text = speed

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_bibliography_doc() -> bytes:
    """Document avec références bibliographiques."""
    doc = Document()
    doc.add_heading("Réseaux de neurones profonds", 0)
    doc.add_heading("Introduction", 1)
    doc.add_paragraph(
        "Les réseaux de neurones convolutifs ont révolutionné la vision par ordinateur "
        "(LeCun et al., 1998). Les architectures modernes comme ResNet (He et al., 2016) "
        "permettent d'entraîner des réseaux très profonds grâce aux connexions résiduelles."
    )
    doc.add_heading("Méthode", 1)
    doc.add_paragraph(
        "L'optimiseur Adam (Kingma & Ba, 2015) est utilisé avec un taux d'apprentissage "
        "de 0,001 et un batch de 32 exemples."
    )
    doc.add_heading("Références", 1)
    refs = [
        "LeCun, Y., Bottou, L., Bengio, Y., & Haffner, P. (1998). "
        "Gradient-based learning applied to document recognition. "
        "Proceedings of the IEEE, 86(11), 2278-2324.",
        "He, K., Zhang, X., Ren, S., & Sun, J. (2016). "
        "Deep residual learning for image recognition. "
        "Proceedings of CVPR, 770-778.",
        "Kingma, D. P., & Ba, J. (2015). "
        "Adam: A method for stochastic optimization. ICLR 2015.",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_special_chars_doc() -> bytes:
    """Caractères spéciaux et accents."""
    doc = Document()
    doc.add_heading("Caractères spéciaux", 1)
    doc.add_paragraph(
        "Français : é, è, ê, à, ù, û, ô, î, ï, ç, œ, æ\n"
        "Allemand : ä, ö, ü, ß\n"
        "Espagnol : ñ, ¿, ¡\n"
        "Symboles : €, ©, ®, ™, °, ±, ×, ÷\n"
        "Mathématiques dans le texte : α, β, γ, δ, π, Σ"
    )
    doc.add_heading("Caractères réservés LaTeX", 2)
    doc.add_paragraph(
        "Ces caractères sont sensibles dans LaTeX : "
        "& % $ # _ { } ~ ^ \\"
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_long_doc() -> bytes:
    """Document avec de nombreuses sections."""
    doc = Document()
    doc.add_heading("Rapport complet de recherche", 0)
    doc.add_paragraph("Auteur : Équipe de recherche\nDate : 2025")

    sections_data = [
        ("Résumé", "Ce rapport présente les résultats d'une étude approfondie sur les "
         "méthodes d'apprentissage automatique appliquées au traitement du langage naturel."),
        ("Introduction", "L'intelligence artificielle connaît un développement sans précédent "
         "depuis la démocratisation des architectures Transformer en 2017."),
        ("État de l'art", "Les modèles de langage pré-entraînés comme BERT, GPT et leurs "
         "variantes ont profondément modifié l'approche du traitement du langage naturel."),
        ("Méthodologie", "Nous avons utilisé un corpus de 10 millions de documents en langue "
         "française, pré-traités avec une pipeline de normalisation textuelle."),
        ("Résultats", "Les expériences montrent une amélioration significative de 15 points "
         "de F1-score par rapport aux baselines établis dans la littérature."),
        ("Discussion", "Ces résultats confirment l'hypothèse initiale selon laquelle "
         "l'augmentation des données d'entraînement améliore les performances."),
        ("Conclusion", "Nous avons démontré l'efficacité des architectures modernes sur "
         "des tâches de classification de texte en français."),
        ("Perspectives", "Les travaux futurs porteront sur l'extension aux langues africaines "
         "et l'optimisation de l'empreinte carbone du modèle."),
    ]

    for title, body in sections_data:
        doc.add_heading(title, 1)
        doc.add_paragraph(body)
        # Sous-sections
        for sub in ["Contexte", "Détails"]:
            doc.add_heading(f"{sub} de {title.lower()}", 2)
            doc.add_paragraph(
                f"Cette section présente les {sub.lower()} concernant {title.lower()}. "
                "Lorem ipsum dolor sit amet, consectetur adipiscing elit. "
                "Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua."
            )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Test runner ────────────────────────────────────────────────────────────────

def assert_zip_contains(data: bytes, filename: str, test_name: str) -> bool:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        if filename not in names:
            fail(f"{test_name} — '{filename}' absent du ZIP (contenu: {names})")
            return False
        ok(f"{test_name} — '{filename}' présent")
        return True


def assert_tex_contains(data: bytes, patterns: list[str], test_name: str):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        tex = z.read("document.tex").decode("utf-8", errors="replace")
    for pat in patterns:
        if pat in tex:
            ok(f"{test_name} — '{pat}' trouvé dans .tex")
        else:
            r = fail(f"{test_name} — '{pat}' ABSENT du .tex")
            warn(r, f"Début du .tex : {tex[:300]}")


def convert(docx_bytes: bytes, filename: str = "test.docx", url: str = BASE_URL):
    return requests.post(
        url + CONVERT,
        files={"file": (filename, docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        timeout=120,
    )


# ── Tests ──────────────────────────────────────────────────────────────────────

def test_health(url: str):
    section("1. Health Check")
    r = requests.get(url + HEALTH, timeout=10)
    if r.status_code == 200 and r.json().get("status") == "ok":
        ok("GET /health", "status=ok")
    else:
        fail("GET /health", f"HTTP {r.status_code} — {r.text[:100]}")


def test_simple(url: str):
    section("2. Document texte simple")
    r = convert(make_simple_doc(), "simple.docx", url)
    if r.status_code != 200:
        fail("Conversion simple", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion simple", f"HTTP {r.status_code}")
    ok("Content-Type zip", r.headers.get("Content-Type", ""))
    ok("Header X-Has-Bibliography", r.headers.get("X-Has-Bibliography", "manquant"))
    ok("Header X-Image-Count", r.headers.get("X-Image-Count", "manquant"))
    assert_zip_contains(r.content, "document.tex", "Simple")
    assert_tex_contains(r.content, [r"\documentclass", r"\begin{document}", r"\end{document}"], "Simple")


def test_formatting(url: str):
    section("3. Formatage (gras, italique, exposant, indice)")
    r = convert(make_rich_formatting_doc(), "formatage.docx", url)
    if r.status_code != 200:
        fail("Conversion formatage", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion formatage", f"HTTP {r.status_code}")
    assert_tex_contains(
        r.content,
        [r"\textbf{", r"\textit{"],
        "Formatage",
    )


def test_lists(url: str):
    section("4. Listes (puces et numéros)")
    r = convert(make_lists_doc(), "listes.docx", url)
    if r.status_code != 200:
        fail("Conversion listes", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion listes", f"HTTP {r.status_code}")
    assert_tex_contains(r.content, [r"\begin{itemize}", r"\begin{enumerate}"], "Listes")


def test_table(url: str):
    section("5. Tableau")
    r = convert(make_table_doc(), "tableau.docx", url)
    if r.status_code != 200:
        fail("Conversion tableau", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion tableau", f"HTTP {r.status_code}")
    assert_tex_contains(r.content, [r"\begin{tabular}", r"\toprule"], "Tableau")


def test_bibliography(url: str):
    section("6. Bibliographie")
    r = convert(make_bibliography_doc(), "biblio.docx", url)
    if r.status_code != 200:
        fail("Conversion biblio", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion biblio", f"HTTP {r.status_code}")
    has_bib = r.headers.get("X-Has-Bibliography") == "true"
    if has_bib:
        ok("X-Has-Bibliography: true")
        assert_zip_contains(r.content, "references.bib", "Biblio")
        assert_tex_contains(r.content, [r"\bibliography{references}"], "Biblio")
    else:
        res = fail("X-Has-Bibliography devrait être true")
        warn(res, "Claude n'a peut-être pas détecté les références — vérifiez manuellement")


def test_special_chars(url: str):
    section("7. Caractères spéciaux")
    r = convert(make_special_chars_doc(), "special.docx", url)
    if r.status_code != 200:
        fail("Conversion spécial", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion spécial", f"HTTP {r.status_code}")
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        tex = z.read("document.tex").decode("utf-8", errors="replace")
    assert_tex_contains(r.content, [r"\usepackage[utf8]{inputenc}"], "Spécial")
    question_marks = tex.count("???")
    if question_marks == 0:
        ok("Spécial — pas de '???' (encodage OK)")
    else:
        fail(f"Spécial — {question_marks} occurrence(s) de '???' (problème encodage)")


def test_long_doc(url: str):
    section("8. Document long (multi-sections)")
    r = convert(make_long_doc(), "long.docx", url)
    if r.status_code != 200:
        fail("Conversion long", f"HTTP {r.status_code} — {r.text[:200]}")
        return
    ok("Conversion long", f"HTTP {r.status_code}")
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        tex = z.read("document.tex").decode("utf-8", errors="replace")
    section_count = tex.count(r"\section{")
    subsection_count = tex.count(r"\subsection{")
    ok(f"Long — {section_count} \\section trouvées")
    ok(f"Long — {subsection_count} \\subsection trouvées")
    if section_count < 3:
        fail("Long — moins de 3 \\section (attendu ≥ 8)")


def test_errors(url: str):
    section("9. Cas d'erreur")

    # Mauvaise extension
    r = requests.post(url + CONVERT, files={"file": ("rapport.pdf", b"%PDF-1.4", "application/pdf")}, timeout=15)
    if r.status_code == 400:
        ok("Erreur E2 — .pdf refusé", f"HTTP 400 : {r.json().get('detail','')}")
    else:
        fail("Erreur E2 — .pdf devrait retourner 400", f"HTTP {r.status_code}")

    # Fichier vide
    r = requests.post(url + CONVERT, files={"file": ("vide.docx", b"", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, timeout=15)
    if r.status_code in (400, 422):
        ok("Erreur E4 — fichier vide refusé", f"HTTP {r.status_code} : {r.json().get('detail','')}")
    else:
        fail("Erreur E4 — fichier vide devrait retourner 400/422", f"HTTP {r.status_code}")

    # Fichier corrompu
    r = requests.post(url + CONVERT, files={"file": ("corrompu.docx", b"PK\x03\x04GARBAGE", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, timeout=15)
    if r.status_code in (422, 400):
        ok("Erreur E5 — fichier corrompu refusé", f"HTTP {r.status_code}")
    else:
        fail("Erreur E5 — corrompu devrait retourner 422", f"HTTP {r.status_code}")

    # Sans fichier
    r = requests.post(url + CONVERT, timeout=15)
    if r.status_code == 422:
        ok("Erreur E1 — sans fichier refusé", "HTTP 422")
    else:
        fail("Erreur E1 — sans fichier devrait retourner 422", f"HTTP {r.status_code}")


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Test Word→LaTeX API")
    parser.add_argument("--url", default="http://localhost:8000", help="URL de base du backend")
    parser.add_argument("--skip-long", action="store_true", help="Ignorer le test du document long")
    args = parser.parse_args()

    url = args.url.rstrip("/")
    print(f"\n{BOLD}🧪 Tests — Word to LaTeX API{RESET}")
    print(f"   Backend : {url}")

    # Vérifier la connectivité
    try:
        requests.get(url + HEALTH, timeout=5)
    except requests.exceptions.ConnectionError:
        print(f"\n{RED}❌ Impossible de se connecter à {url}{RESET}")
        print("   Lancez d'abord : uvicorn main:app --reload --port 8000")
        sys.exit(1)

    test_health(url)
    test_simple(url)
    test_formatting(url)
    test_lists(url)
    test_table(url)
    test_bibliography(url)
    test_special_chars(url)
    if not args.skip_long:
        test_long_doc(url)
    test_errors(url)

    # Résumé
    passed = sum(1 for r in results if r.passed)
    failed = sum(1 for r in results if not r.passed)
    total  = len(results)

    print(f"\n{'═'*55}")
    print(f"{BOLD}  RÉSULTATS : {passed}/{total} tests passés{RESET}")
    if failed:
        print(f"  {RED}{failed} échec(s){RESET}")
        print(f"\n{BOLD}  Échecs :{RESET}")
        for r in results:
            if not r.passed:
                print(f"    {RED}✘{RESET} {r.name}")
                if r.detail:
                    print(f"      {r.detail}")
    else:
        print(f"  {GREEN}Tous les tests sont passés ✔{RESET}")
    print(f"{'═'*55}\n")

    sys.exit(0 if failed == 0 else 1)


if __name__ == "__main__":
    main()
